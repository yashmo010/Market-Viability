"""Apply final targeted fixes strictly to the two READY papers:
1. Research_Paper_IEEE_CORRECTED_v3.docx
2. Survey_Paper_IEEE_CORRECTED_v3.docx
"""

import docx
from pathlib import Path

READY_DIR = Path("Capstone updated/Capstone/submission/ready")
RP_PATH = READY_DIR / "Research_Paper_IEEE_CORRECTED_v3.docx"
SP_PATH = READY_DIR / "Survey_Paper_IEEE_CORRECTED_v3.docx"
NEW_IMG_PATH = Path("scratch_inspect/image22_new.jpg")

def apply_fixes():
    # -------------------------------------------------------------------------
    # 1. RESEARCH PAPER TARGETED FIXES
    # -------------------------------------------------------------------------
    print("Modifying Research Paper...")
    doc_rp = docx.Document(str(RP_PATH))

    # A) Replace Figure 8c screenshot (rId30 -> media/image22.jpg)
    if "rId30" in doc_rp.part.rels and NEW_IMG_PATH.exists():
        new_bytes = NEW_IMG_PATH.read_bytes()
        doc_rp.part.rels["rId30"].target_part._blob = new_bytes
        print(f"  [Issue 1B] Replaced Figure 8c (rId30) with updated app screenshot ({len(new_bytes)} bytes)")
    else:
        print("  [Issue 1B] WARNING: rId30 not found or image missing!")

    # B) Extraction wording clarification in P15
    p15 = doc_rp.paragraphs[15]
    old_p15 = "C1. A null-forcing aspect-extraction scheme over 43,599 reviews:"
    new_p15 = "C1. A null-forcing aspect-extraction scheme across 29,410 extracted reviews (from the 43,599-review corpus):"
    if old_p15 in p15.runs[0].text:
        p15.runs[0].text = p15.runs[0].text.replace(old_p15, new_p15)
        print("  [Issue 4] Updated P15 extraction coverage wording")
    else:
        print("  [Issue 4] WARNING: old_p15 not found in P15!")

    # C) Extraction wording clarification in P37
    p37 = doc_rp.paragraphs[37]
    old_p37_tail = "power banks 99."
    new_p37_tail = "power banks 99. Across the corpus, aspect-level sentiment extraction yielded 29,410 review rows (67.5% review-level coverage, spanning 100% of retained products)."
    if old_p37_tail in p37.runs[0].text:
        p37.runs[0].text = p37.runs[0].text.replace(old_p37_tail, new_p37_tail)
        print("  [Issue 4] Updated P37 extraction coverage wording")
    else:
        print("  [Issue 4] WARNING: old_p37_tail not found in P37!")

    # D) Remove unsupported brand-tier baseline sentence in P84 (Section V.C)
    p84 = doc_rp.paragraphs[84]
    unsupported_sentence = "On the small n = 59 frozen holdout a simple brand-tier baseline (ρ ≈ 0.41 in our re-analysis) ranks products about as well as the model, and it is itself subject to the full-catalog leakage caveat above. "
    if unsupported_sentence in p84.runs[0].text:
        p84.runs[0].text = p84.runs[0].text.replace(unsupported_sentence, "")
        print("  [Issue 3] Removed unsupported brand-tier baseline sentence from Section V.C (P84)")
    else:
        print("  [Issue 3] WARNING: unsupported_sentence not found in P84!")

    doc_rp.save(str(RP_PATH))
    print("Saved Research Paper.")

    # -------------------------------------------------------------------------
    # 2. SURVEY PAPER TARGETED FIXES
    # -------------------------------------------------------------------------
    print("\nModifying Survey Paper...")
    doc_sp = docx.Document(str(SP_PATH))

    # A) Extraction wording in P15
    p15_sp = doc_sp.paragraphs[15]
    found_p15_sp = False
    for r in p15_sp.runs:
        if "from 43,599 historical customer" in r.text:
            r.text = r.text.replace("from 43,599 historical customer", "across 29,410 reviews from a 43,599-review corpus of")
            found_p15_sp = True
            break
    if found_p15_sp:
        print("  [Issue 4] Updated Survey Paper P15 extraction coverage wording")
    else:
        print("  [Issue 4] WARNING: target text not found in Survey P15!")

    # B) Extraction wording in P42
    p42_sp = doc_sp.paragraphs[42]
    old_p42_cat = "(Bluetooth speakers, ice makers, power banks, smartwatches, and wireless headphones)."
    new_p42_cat = "(Bluetooth speakers, ice makers, power banks, smartwatches, and wireless headphones), with 29,410 reviews possessing extracted aspect representations across all 631 retained physical products."
    found_p42_sp = False
    for r in p42_sp.runs:
        if old_p42_cat in r.text:
            r.text = r.text.replace(old_p42_cat, new_p42_cat)
            found_p42_sp = True
            break
    if found_p42_sp:
        print("  [Issue 4] Updated Survey Paper P42 extraction coverage wording")
    else:
        print("  [Issue 4] WARNING: target text not found in Survey P42!")

    doc_sp.save(str(SP_PATH))
    print("Saved Survey Paper.")

if __name__ == "__main__":
    apply_fixes()
