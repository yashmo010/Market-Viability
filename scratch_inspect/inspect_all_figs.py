import docx
import xml.etree.ElementTree as ET

doc = docx.Document(r'redesign\Research_Paper_IEEE_Final.docx')
for i in range(60, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    rids = []
    for r in p.runs:
        if 'blip' in r._element.xml:
            tree = ET.fromstring(r._element.xml)
            for elem in tree.iter():
                if elem.tag.endswith('blip'):
                    rids.append(elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'))
    if rids or 'Figure' in p.text:
        img_names = [doc.part.rels[rid].target_ref for rid in rids if rid]
        safe_text = p.text[:80].encode('ascii', 'replace').decode('ascii')
        print(f"P{i}: text=\"{safe_text}\" images={img_names}")
