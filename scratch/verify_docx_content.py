import docx
from pathlib import Path

p = Path(r"c:\Users\LENOVO\OneDrive\Desktop\Capstone updated (3)\Legitimate\09_MANUSCRIPT\FINAL_MANUSCRIPT.docx")
doc = docx.Document(str(p))

print(f"Total paragraphs in FINAL_MANUSCRIPT.docx: {len(doc.paragraphs)}")
for i, para in enumerate(doc.paragraphs[:25]):
    txt = para.text.strip()
    if txt:
        print(f"[{i}] {txt[:120]}...")
