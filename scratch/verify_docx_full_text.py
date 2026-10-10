import docx
from pathlib import Path

p = Path(r"c:\Users\LENOVO\OneDrive\Desktop\Capstone updated (3)\Legitimate\09_MANUSCRIPT\FINAL_MANUSCRIPT.docx")
doc = docx.Document(str(p))

print("=== Section II Text ===")
print(doc.paragraphs[13].text)
print("\n" + doc.paragraphs[14].text)
print("\n=== Figure 2 Caption ===")
print(doc.paragraphs[16].text)

for i, para in enumerate(doc.paragraphs):
    if "VI. EXTERNAL" in para.text or "VII. DISCUSSION" in para.text:
        print(f"\n=== Paragraph {i}: {para.text} ===")
        if i + 1 < len(doc.paragraphs):
            print(doc.paragraphs[i+1].text)
        if i + 2 < len(doc.paragraphs):
            print(doc.paragraphs[i+2].text)
